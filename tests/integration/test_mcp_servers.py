"""Real stdio protocol tests with mock external services and real report files."""

from __future__ import annotations

import asyncio
import json
import os
import sys
from contextlib import asynccontextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from mcp.types.version import LATEST_HANDSHAKE_VERSION
from pptx import Presentation

pytestmark = pytest.mark.integration

_ROOT = Path(__file__).resolve().parents[2]
_SERVERS = {
    "researcher": ("researcher-agent", ["research_company"]),
    "sharepoint": ("sharepoint-agent", ["search_documents", "get_document_content"]),
    "report_generator": ("report-generator", ["generate_report"]),
}
_VALID_ARGUMENTS = {
    "research_company": {"company_name": "Tailspin Toys"},
    "search_documents": {"query": "Tailspin"},
    "get_document_content": {"drive_id": "sample-drive", "item_id": "sample-item"},
    "generate_report": {"title": "MCP Sales Report", "customer_name": "Tailspin Toys"},
}
_SCHEMA_CONTRACTS = {
    "research_company": {
        "type": "object",
        "properties": {"company_name": {"type": "string"}, "focus_areas": {"type": "string"}},
        "required": ["company_name"],
    },
    "search_documents": {
        "type": "object",
        "properties": {"query": {"type": "string"}, "site_id": {"type": "string"}},
        "required": ["query"],
    },
    "get_document_content": {
        "type": "object",
        "properties": {"drive_id": {"type": "string"}, "item_id": {"type": "string"}},
        "required": ["drive_id", "item_id"],
    },
    "generate_report": {
        "type": "object",
        "properties": {
            "title": {"type": "string"},
            "customer_name": {"type": "string"},
            "format": {"type": "string", "enum": ["docx", "pptx"], "default": "docx"},
            "pipeline_data": {"type": "array", "items": {"type": "object"}},
            "research_data": {"type": "object", "additionalProperties": True},
            "sharepoint_docs": {"type": "array", "items": {"type": "object"}},
            "forecast_data": {"type": "object", "additionalProperties": True},
            "additional_context": {"type": "string"},
        },
        "required": ["title", "customer_name"],
        "additionalProperties": False,
    },
}


def _without_descriptions(value):
    if isinstance(value, dict):
        return {key: _without_descriptions(item) for key, item in value.items() if key != "description"}
    if isinstance(value, list):
        return [_without_descriptions(item) for item in value]
    return value


@dataclass
class StdioClient:
    process: asyncio.subprocess.Process
    stderr: list[str] = field(default_factory=list)
    notifications: list[dict[str, Any]] = field(default_factory=list)
    request_id: int = 0

    async def send(self, message: dict[str, Any]) -> None:
        assert self.process.stdin is not None
        self.process.stdin.write((json.dumps(message) + "\n").encode())
        await self.process.stdin.drain()

    async def request(self, method: str, params: Any = None) -> dict[str, Any]:
        self.request_id += 1
        message = {"jsonrpc": "2.0", "id": self.request_id, "method": method}
        if params is not None:
            message["params"] = params
        await self.send(message)
        assert self.process.stdout is not None
        async with asyncio.timeout(15):
            while True:
                line = await self.process.stdout.readline()
                assert line, f"Server closed stdout: {''.join(self.stderr)}"
                response = json.loads(line)
                assert response["jsonrpc"] == "2.0"
                if response.get("id") == self.request_id:
                    return response
                assert response.get("id") is None, response
                self.notifications.append(response)

    async def call(self, name: str, arguments: Any) -> dict[str, Any]:
        return await self.request("tools/call", {"name": name, "arguments": arguments})


@asynccontextmanager
async def _server(server_name: str, cwd: Path, protocol: str = "2024-11-05", fault: str | None = None):
    module = f"src.agents.{server_name}.mcp_server"
    command = ["-m", module]
    if fault is not None:
        target = {
            "researcher": "research_company",
            "sharepoint": "search_documents",
            "report_generator": "generate_docx",
        }[server_name]
        failure = "RuntimeError('programmer fault')"
        if fault == "operational":
            failure = {
                "researcher": "json.JSONDecodeError('invalid provider JSON', '', 0)",
                "sharepoint": "ClientAuthenticationError('credential unavailable')",
                "report_generator": "PermissionError('cannot write report')",
            }[server_name]
        definition = "def" if server_name == "report_generator" else "async def"
        script = (
            "import asyncio, importlib, json, sys\n"
            "from azure.core.exceptions import ClientAuthenticationError\n"
            "module = importlib.import_module(sys.argv[1])\n"
            f"{definition} fail(*args, **kwargs):\n"
            f"    raise {failure}\n"
            f"module.{target} = fail\n"
            "asyncio.run(module.main())\n"
        )
        command = ["-c", script, module]

    process = await asyncio.create_subprocess_exec(
        sys.executable,
        *command,
        cwd=cwd,
        env={
            **os.environ,
            "PYTHONPATH": os.pathsep.join(filter(None, [str(_ROOT), os.environ.get("PYTHONPATH")])),
            "SEARCH_PROVIDER": "mock",
            "SHAREPOINT_MODE": "mock",
        },
        stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.PIPE,
    )
    client = StdioClient(process)

    async def drain_stderr():
        assert process.stderr is not None
        async for line in process.stderr:
            client.stderr.append(line.decode(errors="replace"))

    stderr_task = asyncio.create_task(drain_stderr())
    try:
        initialized = await client.request(
            "initialize",
            {
                "protocolVersion": protocol,
                "capabilities": {},
                "clientInfo": {"name": "test-client", "version": "0.1.0"},
            },
        )
        assert "error" not in initialized, initialized
        result = initialized["result"]
        assert result["serverInfo"]["name"] == _SERVERS[server_name][0]
        assert result["protocolVersion"] == protocol
        assert "tools" in result["capabilities"]
        await client.send({"jsonrpc": "2.0", "method": "notifications/initialized"})
        yield client
    finally:
        if process.returncode is None:
            process.terminate()
            try:
                await asyncio.wait_for(process.wait(), timeout=5)
            except TimeoutError:
                process.kill()
                await process.wait()
        await stderr_task


def _tool_result(response, *, is_error=False):
    assert "error" not in response, response
    result = response["result"]
    assert result["isError"] is is_error
    assert "is_error" not in result
    assert "structuredContent" not in result
    assert len(result["content"]) == 1
    assert result["content"][0]["type"] == "text"
    return result["content"][0]["text"]


@pytest.mark.parametrize("server_name", _SERVERS)
@pytest.mark.parametrize("protocol", ["2024-11-05", LATEST_HANDSHAKE_VERSION])
async def test_initialize_and_advertised_schema_contract(server_name, protocol, tmp_path):
    async with _server(server_name, tmp_path, protocol) as client:
        response = await client.request("tools/list", {})
        assert "error" not in response, response
        tools = response["result"]["tools"]
        assert [tool["name"] for tool in tools] == _SERVERS[server_name][1]
        for tool in tools:
            assert "input_schema" not in tool
            assert "outputSchema" not in tool
            assert _without_descriptions(tool["inputSchema"]) == _SCHEMA_CONTRACTS[tool["name"]]


async def test_research_valid_calls_before_listing(tmp_path):
    async with _server("researcher", tmp_path) as client:
        text = _tool_result(
            await client.call(
                "research_company", {"company_name": "Tailspin Toys", "focus_areas": "custom", "extra": 1}
            )
        )
        data = json.loads(text)
        assert data["company_name"] == "Tailspin Toys"
        assert data["articles"]
        assert data["key_metrics"]
        generic = json.loads(_tool_result(await client.call("research_company", {"company_name": "UnknownCorp"})))
        assert generic["company_name"] == "UnknownCorp"
        assert generic["articles"] == []


async def test_sharepoint_valid_calls_before_listing(tmp_path):
    async with _server("sharepoint", tmp_path) as client:
        documents = json.loads(
            _tool_result(await client.call("search_documents", {"query": "Tailspin", "site_id": "sample", "extra": 1}))
        )
        assert documents
        assert "Tailspin" in documents[0]["name"]
        assert documents[0]["url"]
        empty = json.loads(_tool_result(await client.call("search_documents", {"query": "nonexistent-xyz-12345"})))
        assert empty == []
        content = json.loads(
            _tool_result(
                await client.call(
                    "get_document_content", {"drive_id": "sample-drive", "item_id": "sample-item", "extra": 1}
                )
            )
        )
        assert "Tailspin Toys" in content["content_text"]
        assert content["url"]
        assert content["size"] > 0


async def test_report_valid_calls_create_readable_artifacts_before_listing(tmp_path):
    async with _server("report_generator", tmp_path) as client:
        for format_name in ["docx", "pptx"]:
            arguments = {
                **_VALID_ARGUMENTS["generate_report"],
                "research_data": {
                    "articles": [{"title": "Sample research citation", "url": "https://example.com/research"}]
                },
                "sharepoint_docs": [
                    {"name": "Sample account plan", "url": "https://example.com/plan", "excerpt": "Sample context"}
                ],
            }
            if format_name == "pptx":
                arguments["format"] = format_name
            data = json.loads(_tool_result(await client.call("generate_report", arguments)))
            assert data["status"] == "success"
            assert data["format"] == format_name
            assert data["title"] == arguments["title"]
            output = tmp_path / data["file_path"]
            assert output.resolve().is_relative_to(tmp_path)
            assert output.suffix == f".{format_name}"
            if format_name == "docx":
                document = Document(output)
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                assert "Pipeline Overview" in text
                assert "Sources & Citations" in text
            else:
                presentation = Presentation(output)
                text = "\n".join(
                    shape.text_frame.text
                    for slide in presentation.slides
                    for shape in slide.shapes
                    if shape.has_text_frame
                )
                assert len(presentation.slides) >= 4
            assert arguments["title"] in text
            assert "Sample research citation" in text
            assert "Sample account plan" in text


@pytest.mark.parametrize("server_name", _SERVERS)
async def test_schema_rejections_and_recovery(server_name, tmp_path):
    async with _server(server_name, tmp_path) as client:
        for tool_name in _SERVERS[server_name][1]:
            valid = _VALID_ARGUMENTS[tool_name]
            schema = _SCHEMA_CONTRACTS[tool_name]
            rejected = [None, {}]
            rejected.extend(
                {key: value for key, value in valid.items() if key != required} for required in schema["required"]
            )
            invalid_values = {
                "string": [None, True, 42, [], {}],
                "array": [None, True, 42, "wrong", {}],
                "object": [None, True, 42, "wrong", []],
            }
            for key, property_schema in schema["properties"].items():
                rejected.extend({**valid, key: value} for value in invalid_values[property_schema["type"]])
            if tool_name == "generate_report":
                rejected.extend({**valid, "format": value} for value in ["pdf", "DOCX", "PPTX", ""])
                rejected.append({**valid, "unknown": "extra"})
                rejected.extend(
                    {**valid, field: [value]}
                    for field in ["pipeline_data", "sharepoint_docs"]
                    for value in [None, True, 42, "wrong", []]
                )
            for arguments in rejected:
                text = _tool_result(await client.call(tool_name, arguments), is_error=True)
                assert text.startswith("Input validation error: ")
                assert not (tmp_path / "output").exists()
            missing = await client.request("tools/call", {"name": tool_name})
            assert _tool_result(missing, is_error=True).startswith("Input validation error: ")
        unknown = await client.call("unknown_tool", {})
        assert _tool_result(unknown, is_error=True) == "Unknown tool: unknown_tool"
        assert not (tmp_path / "output").exists()
        for tool_name in _SERVERS[server_name][1]:
            _tool_result(await client.call(tool_name, _VALID_ARGUMENTS[tool_name]))


@pytest.mark.parametrize("server_name", _SERVERS)
async def test_protocol_rejections_and_recovery(server_name, tmp_path):
    async with _server(server_name, tmp_path) as client:
        tool_name = _SERVERS[server_name][1][0]
        for arguments in [[], "wrong", 42, True]:
            response = await client.call(tool_name, arguments)
            assert "result" not in response
            assert response["error"]["code"] == -32602
        for params in [{"arguments": {}}, {"name": 42, "arguments": {}}]:
            response = await client.request("tools/call", params)
            assert "result" not in response
            assert response["error"]["code"] == -32602
        unknown = await client.request("nonexistent/method", {})
        assert unknown["error"]["code"] == -32601
        assert not (tmp_path / "output").exists()
        # Invalid envelopes are rejected by the transport before request dispatch.
        await client.send({"jsonrpc": "2.0", "id": "invalid-envelope", "method": "tools/call", "params": []})
        assert client.process.stdin is not None
        client.process.stdin.write(b"{not valid json}\n")
        await client.process.stdin.drain()
        ping = await client.request("ping", {})
        assert "result" in ping
        assert "error" not in ping
        assert not (tmp_path / "output").exists()
        _tool_result(await client.call(tool_name, _VALID_ARGUMENTS[tool_name]))


@pytest.mark.parametrize("server_name", _SERVERS)
@pytest.mark.parametrize("fault", ["operational", "unexpected"])
async def test_error_envelopes_from_real_dispatch(server_name, fault, tmp_path):
    async with _server(server_name, tmp_path, fault=fault) as client:
        tool_name = _SERVERS[server_name][1][0]
        response = await client.call(tool_name, _VALID_ARGUMENTS[tool_name])
        if fault == "operational":
            assert _tool_result(response, is_error=True)
        else:
            assert "result" not in response
            assert response["error"]["code"] == 0
            assert response["error"]["message"] == "programmer fault"
        assert "result" in await client.request("ping", {})
