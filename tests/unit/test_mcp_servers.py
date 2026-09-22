"""Validation and error contracts at the MCP callback boundary."""

from __future__ import annotations

import asyncio
import copy
import io
import json
from unittest.mock import AsyncMock, Mock
from zipfile import ZipFile

import pytest
from aiohttp import ClientConnectionError, ClientResponseError
from azure.core.exceptions import ClientAuthenticationError
from docx import Document
from mcp.types import CallToolRequestParams, TextContent

from src.agents.mcp_validation import validate_tool_arguments
from src.agents.report_generator import mcp_server as report
from src.agents.researcher import mcp_server as researcher
from src.agents.sharepoint import mcp_server as sharepoint
from src.agents.sharepoint import tools as sharepoint_tools

_CASES = [
    (researcher.handle_call_tool, researcher.handle_list_tools, "research_company", {"company_name": "Tailspin Toys"}),
    (sharepoint.call_tool, sharepoint.list_tools, "search_documents", {"query": "Tailspin"}),
    (
        sharepoint.call_tool,
        sharepoint.list_tools,
        "get_document_content",
        {"drive_id": "sample-drive", "item_id": "sample-item"},
    ),
    (
        report.call_tool,
        report.list_tools,
        "generate_report",
        {"title": "Sales Report", "customer_name": "Tailspin Toys"},
    ),
]


@pytest.fixture(params=_CASES, ids=[case[2] for case in _CASES])
def tool_case(request):
    return request.param


@pytest.fixture
def no_side_effects(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    mocks = [
        (researcher, "research_company", AsyncMock()),
        (sharepoint, "search_documents", AsyncMock()),
        (sharepoint, "get_document_content", AsyncMock()),
        (report, "_build_report_data", Mock()),
        (report, "generate_docx", Mock()),
        (report, "generate_pptx", Mock()),
        (report.Path, "mkdir", Mock()),
    ]
    for module, name, mock in mocks:
        mock.side_effect = AssertionError(f"Rejected input reached {name}")
        monkeypatch.setattr(module, name, mock)
    yield
    for _, _, mock in mocks:
        mock.assert_not_called()
    assert not (tmp_path / "output").exists()


def _assert_validation_error(result):
    assert result.is_error is True
    assert result.structured_content is None
    assert len(result.content) == 1
    assert isinstance(result.content[0], TextContent)
    assert result.content[0].text.startswith("Input validation error: ")


async def test_missing_required_fields(tool_case, no_side_effects):
    call, listing, name, valid = tool_case
    tool = next(tool for tool in (await listing(Mock(), None)).tools if tool.name == name)
    for field in tool.input_schema["required"]:
        arguments = {key: value for key, value in valid.items() if key != field}
        _assert_validation_error(await call(Mock(), CallToolRequestParams(name=name, arguments=arguments)))


@pytest.mark.parametrize("arguments", [None, {}])
async def test_missing_arguments(tool_case, arguments, no_side_effects):
    call, _, name, _ = tool_case
    _assert_validation_error(await call(Mock(), CallToolRequestParams(name=name, arguments=arguments)))


async def test_wrong_types_for_every_advertised_field(tool_case, no_side_effects):
    call, listing, name, valid = tool_case
    tool = next(tool for tool in (await listing(Mock(), None)).tools if tool.name == name)
    invalid_values = {
        "string": [None, True, 42, 1.5, [], {}],
        "array": [None, True, 42, "not-an-array", {}],
        "object": [None, True, 42, "not-an-object", []],
    }
    for field, schema in tool.input_schema["properties"].items():
        for value in invalid_values[schema["type"]]:
            arguments = {**valid, field: value}
            _assert_validation_error(await call(Mock(), CallToolRequestParams(name=name, arguments=arguments)))


@pytest.mark.parametrize("field", ["pipeline_data", "sharepoint_docs"])
@pytest.mark.parametrize("item", [None, True, 42, "not-an-object", []])
async def test_report_array_items(field, item, no_side_effects):
    arguments = {"title": "Report", "customer_name": "Sample", field: [item]}
    _assert_validation_error(
        await report.call_tool(Mock(), CallToolRequestParams(name="generate_report", arguments=arguments))
    )


@pytest.mark.parametrize("format_name", ["pdf", "DOCX", "PPTX", ""])
async def test_report_format_is_exact_enum(format_name, no_side_effects):
    arguments = {"title": "Report", "customer_name": "Sample", "format": format_name}
    _assert_validation_error(
        await report.call_tool(Mock(), CallToolRequestParams(name="generate_report", arguments=arguments))
    )


async def test_report_rejects_extra_fields(no_side_effects):
    arguments = {"title": "Report", "customer_name": "Sample", "sections": []}
    _assert_validation_error(
        await report.call_tool(Mock(), CallToolRequestParams(name="generate_report", arguments=arguments))
    )


async def test_unknown_tool_is_tool_error(tool_case, no_side_effects):
    call, _, _, _ = tool_case
    result = await call(Mock(), CallToolRequestParams(name="unknown_tool", arguments={}))
    assert result.is_error is True
    assert result.content == [TextContent(type="text", text="Unknown tool: unknown_tool")]


async def test_advertised_schema_is_validation_source(tool_case):
    _, listing, name, valid = tool_case
    tool = next(tool for tool in (await listing(Mock(), None)).tools if tool.name == name)
    schema_before = copy.deepcopy(tool.input_schema)
    arguments = dict(valid)
    assert validate_tool_arguments(tool, arguments) is None
    assert arguments == valid
    assert tool.input_schema == schema_before
    assert tool.model_dump(by_alias=True)["inputSchema"] == schema_before
    for required in schema_before["required"]:
        assert validate_tool_arguments(tool, {**valid, required: ""}) is None


async def test_report_nested_objects_remain_permissive():
    tool = (await report.list_tools(Mock(), None)).tools[0]
    arguments = {
        "title": "Report",
        "customer_name": "Sample",
        "pipeline_data": [{"custom": {"value": True}}],
        "sharepoint_docs": [{"custom": []}],
        "research_data": {"custom": None},
        "forecast_data": {"custom": 42},
    }
    assert validate_tool_arguments(tool, arguments) is None
    assert "format" not in arguments


async def test_research_preserves_extra_keys_and_unrestricted_focus(monkeypatch):
    handler = AsyncMock(return_value={"company_name": "Sample", "articles": []})
    monkeypatch.setattr(researcher, "research_company", handler)
    result = await researcher.handle_call_tool(
        Mock(),
        CallToolRequestParams(
            name="research_company",
            arguments={"company_name": "Sample", "focus_areas": "custom focus", "extra": True},
        ),
    )
    assert result.is_error is False
    handler.assert_awaited_once_with(company_name="Sample", focus_areas="custom focus")
    assert json.loads(result.content[0].text) == handler.return_value


@pytest.mark.parametrize("case", _CASES[1:3], ids=["search", "content"])
async def test_sharepoint_preserves_extra_keys(monkeypatch, case):
    call, _, name, valid = case
    handler = AsyncMock(return_value={})
    monkeypatch.setattr(sharepoint, name, handler)
    result = await call(Mock(), CallToolRequestParams(name=name, arguments={**valid, "extra": 42}))
    assert result.is_error is False
    if name == "search_documents":
        handler.assert_awaited_once_with(valid["query"], site_id=None)
    else:
        handler.assert_awaited_once_with(valid["drive_id"], valid["item_id"])


@pytest.mark.parametrize(
    "failure",
    [
        ClientAuthenticationError("credential unavailable"),
        ClientConnectionError("connection unavailable"),
        ClientResponseError(Mock(real_url="https://example.invalid"), (), status=403, message="forbidden"),
        TimeoutError("request timed out"),
        json.JSONDecodeError("invalid response", "", 0),
        sharepoint_tools.DocumentContentError("invalid document"),
    ],
)
async def test_sharepoint_operational_failures_are_tool_errors(monkeypatch, failure):
    handler = AsyncMock(side_effect=failure)
    monkeypatch.setattr(sharepoint, "get_document_content", handler)
    result = await sharepoint.call_tool(
        Mock(),
        CallToolRequestParams(name="get_document_content", arguments={"drive_id": "drive", "item_id": "item"}),
    )
    assert result.is_error is True
    assert result.content == [TextContent(type="text", text=str(failure))]


@pytest.mark.parametrize(
    "document_kind", ["invalid-zip", "missing-parts", "invalid-xml", "wrong-content-type", "valid"]
)
async def test_sharepoint_document_parsing_boundary(monkeypatch, document_kind):
    if document_kind == "invalid-zip":
        raw = b"not a zip package"
    else:
        buffer = io.BytesIO()
        if document_kind == "valid":
            document = Document()
            document.add_paragraph("Sample document content")
            document.save(buffer)
        elif document_kind == "wrong-content-type":
            original = io.BytesIO()
            Document().save(original)
            with ZipFile(original) as source, ZipFile(buffer, "w") as destination:
                for name in source.namelist():
                    content = source.read(name)
                    if name == "[Content_Types].xml":
                        content = content.replace(
                            b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                        )
                    destination.writestr(name, content)
        else:
            with ZipFile(buffer, "w") as archive:
                if document_kind == "invalid-xml":
                    archive.writestr("[Content_Types].xml", b"<malformed")
                else:
                    archive.writestr("unrelated.txt", b"missing document parts")
        raw = buffer.getvalue()

    graph = AsyncMock(
        side_effect=[
            {"name": "Sample.docx", "webUrl": "https://example.com/sample"},
            {
                "_raw_bytes": raw,
                "_content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            },
        ]
    )
    monkeypatch.setenv("SHAREPOINT_MODE", "graph")
    monkeypatch.setattr(sharepoint_tools, "_graph_request", graph)
    result = await sharepoint.call_tool(
        Mock(), CallToolRequestParams(name="get_document_content", arguments={"drive_id": "drive", "item_id": "item"})
    )
    assert graph.await_count == 2
    assert result.is_error is (document_kind != "valid")
    if document_kind == "valid":
        assert json.loads(result.content[0].text)["content_text"] == "Sample document content"
    else:
        assert result.content[0].text
        assert result.structured_content is None


async def test_research_decode_failure_is_tool_error(monkeypatch):
    failure = json.JSONDecodeError("invalid response", "", 0)
    monkeypatch.setattr(researcher, "research_company", AsyncMock(side_effect=failure))
    result = await researcher.handle_call_tool(
        Mock(), CallToolRequestParams(name="research_company", arguments={"company_name": "Sample"})
    )
    assert result.is_error is True
    assert result.content == [TextContent(type="text", text=str(failure))]


@pytest.mark.parametrize("format_name", ["docx", "pptx"])
@pytest.mark.parametrize(
    "failure",
    [PermissionError("cannot write report"), UnicodeDecodeError("utf-8", b"\xff", 0, 1, "invalid template")],
)
async def test_report_operational_failures_are_tool_errors(monkeypatch, tmp_path, format_name, failure):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(report, f"generate_{format_name}", Mock(side_effect=failure))
    result = await report.call_tool(
        Mock(),
        CallToolRequestParams(
            name="generate_report", arguments={"title": "Report", "customer_name": "Sample", "format": format_name}
        ),
    )
    assert result.is_error is True
    assert result.content == [TextContent(type="text", text=str(failure))]
    assert not list((tmp_path / "output").iterdir())


@pytest.mark.parametrize(
    "failure",
    [
        RuntimeError("programmer fault"),
        TypeError("programmer fault"),
        KeyError("programmer fault"),
        ValueError("programmer fault"),
        asyncio.CancelledError(),
    ],
)
async def test_unexpected_failures_and_cancellation_propagate(tool_case, monkeypatch, tmp_path, failure):
    call, _, name, valid = tool_case
    monkeypatch.chdir(tmp_path)
    if name == "research_company":
        monkeypatch.setattr(researcher, name, AsyncMock(side_effect=failure))
    elif name in {"search_documents", "get_document_content"}:
        monkeypatch.setattr(sharepoint, name, AsyncMock(side_effect=failure))
    else:
        monkeypatch.setattr(report, "generate_docx", Mock(side_effect=failure))
    with pytest.raises(type(failure)):
        await call(Mock(), CallToolRequestParams(name=name, arguments=valid))
